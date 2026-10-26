"""Разбор расписания из текста, CSV и Excel.

Формат строки текста:
    Пн 09:00-10:30 Математика | ауд. 101 | Иванов И.И.
Аудитория и преподаватель необязательны. Строка, состоящая только из дня
недели (например, «Понедельник»), задаёт день для следующих строк без дня.

Календарь .ics (экспорт из Google Календаря, Outlook, сайта вуза) сворачивается
в недельный шаблон: каждое событие даёт пару в свой день недели и время.
"""
import csv
import io
import re
from dataclasses import dataclass
from datetime import datetime, time, timedelta, timezone
from zoneinfo import ZoneInfo, ZoneInfoNotFoundError

WEEKDAYS = {
    0: ("пн", "пон", "понедельник", "mon", "monday"),
    1: ("вт", "вто", "вторник", "tue", "tuesday"),
    2: ("ср", "сре", "среда", "wed", "wednesday"),
    3: ("чт", "чет", "четверг", "thu", "thursday"),
    4: ("пт", "пят", "пятница", "fri", "friday"),
    5: ("сб", "суб", "суббота", "sat", "saturday"),
    6: ("вс", "вос", "воскресенье", "sun", "sunday"),
}
WEEKDAY_NAMES = ["Понедельник", "Вторник", "Среда", "Четверг", "Пятница", "Суббота", "Воскресенье"]
WEEKDAY_SHORT = ["Пн", "Вт", "Ср", "Чт", "Пт", "Сб", "Вс"]

_TIME_RANGE = re.compile(r"(\d{1,2})[:.](\d{2})\s*[-–—]\s*(\d{1,2})[:.](\d{2})")


@dataclass
class Lesson:
    weekday: int
    start: time
    end: time
    subject: str
    room: str = ""
    teacher: str = ""


class ScheduleParseError(ValueError):
    pass


def parse_weekday(token: str) -> int | None:
    token = token.strip().lower().rstrip(".,:")
    for day, names in WEEKDAYS.items():
        if token in names:
            return day
    return None


def _parse_time_range(text: str) -> tuple[time, time, str] | None:
    m = _TIME_RANGE.search(text)
    if not m:
        return None
    h1, m1, h2, m2 = map(int, m.groups())
    try:
        start, end = time(h1, m1), time(h2, m2)
    except ValueError:
        return None
    rest = (text[: m.start()] + " " + text[m.end():]).strip()
    return start, end, rest


def parse_text(text: str) -> list[Lesson]:
    lessons: list[Lesson] = []
    errors: list[str] = []
    current_day: int | None = None
    for lineno, raw in enumerate(text.splitlines(), 1):
        line = raw.strip()
        if not line:
            continue
        first, _, remainder = line.partition(" ")
        day = parse_weekday(first)
        if day is not None:
            current_day = day
            line = remainder.strip()
            if not line:
                continue
        if current_day is None:
            errors.append(f"строка {lineno}: не указан день недели")
            continue
        parsed = _parse_time_range(line)
        if not parsed:
            errors.append(f"строка {lineno}: не нашёл время вида 09:00-10:30")
            continue
        start, end, rest = parsed
        parts = [p.strip() for p in rest.split("|")]
        subject = parts[0] if parts else ""
        if not subject:
            errors.append(f"строка {lineno}: не указан предмет")
            continue
        lessons.append(Lesson(
            weekday=current_day,
            start=start,
            end=end,
            subject=subject,
            room=parts[1] if len(parts) > 1 else "",
            teacher=parts[2] if len(parts) > 2 else "",
        ))
    if errors:
        raise ScheduleParseError("\n".join(errors))
    if not lessons:
        raise ScheduleParseError("не нашёл ни одной пары")
    return lessons


def _rows_to_lessons(rows: list[list[str]]) -> list[Lesson]:
    """Строки таблицы: день, время (09:00-10:30), предмет, аудитория, преподаватель.
    Также допускается: день, начало, конец, предмет, аудитория, преподаватель."""
    lines = []
    for row in rows:
        cells = [str(c).strip() if c is not None else "" for c in row]
        if not any(cells):
            continue
        if parse_weekday(cells[0]) is None and lines == [] and cells[0]:
            continue  # заголовок таблицы
        if len(cells) >= 3 and _parse_time_range(cells[1]) is None and _parse_time_range(f"{cells[1]}-{cells[2]}"):
            cells = [cells[0], f"{cells[1]}-{cells[2]}", *cells[3:]]
        day, tr, *rest = cells + [""] * (3 - len(cells))
        extra = " | ".join(rest).rstrip(" |")
        lines.append(f"{day} {tr} {extra}".strip())
    return parse_text("\n".join(lines))


def parse_csv(data: bytes) -> list[Lesson]:
    text = data.decode("utf-8-sig", errors="replace")
    dialect = csv.Sniffer().sniff(text.splitlines()[0]) if text.strip() else csv.excel
    return _rows_to_lessons(list(csv.reader(io.StringIO(text), dialect)))


def _cell_to_str(value) -> str:
    if isinstance(value, time):
        return value.strftime("%H:%M")
    return "" if value is None else str(value)


def parse_xlsx(data: bytes) -> list[Lesson]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.active
    rows = [[_cell_to_str(v) for v in row] for row in ws.iter_rows(values_only=True)]
    return _rows_to_lessons(rows)


_DURATION = re.compile(r"P(?:(\d+)D)?(?:T(?:(\d+)H)?(?:(\d+)M)?(?:(\d+)S)?)?")


def _ics_unescape(value: str) -> str:
    return (value.replace("\\n", " ").replace("\\N", " ").replace("\\,", ",")
            .replace("\\;", ";").replace("\\\\", "\\").strip())


def _ics_events(text: str) -> list[dict[str, tuple[list[str], str]]]:
    lines: list[str] = []
    for raw in text.splitlines():
        if raw[:1] in (" ", "\t") and lines:
            lines[-1] += raw[1:]  # перенос длинной строки
        else:
            lines.append(raw.rstrip("\r"))
    events, current = [], None
    for line in lines:
        upper = line.upper()
        if upper == "BEGIN:VEVENT":
            current = {}
        elif upper == "END:VEVENT":
            if current is not None:
                events.append(current)
            current = None
        elif current is not None and ":" in line:
            head, value = line.split(":", 1)
            name, *params = head.split(";")
            current.setdefault(name.upper(), (params, value))
    return events


def _ics_datetime(params: list[str], value: str, user_tz: ZoneInfo) -> datetime | None:
    value = value.strip()
    if len(value) == 8 or any(p.upper() == "VALUE=DATE" for p in params):
        return None  # событие на весь день, не пара
    try:
        dt = datetime.strptime(value.rstrip("Zz"), "%Y%m%dT%H%M%S")
    except ValueError:
        return None
    if value.upper().endswith("Z"):
        dt = dt.replace(tzinfo=timezone.utc)
    else:
        zone = user_tz
        for p in params:
            if p.upper().startswith("TZID="):
                try:
                    zone = ZoneInfo(p[5:].strip('"'))
                except (ZoneInfoNotFoundError, ValueError):
                    pass  # например, Windows-названия зон: считаем, что это пояс пользователя
        dt = dt.replace(tzinfo=zone)
    return dt.astimezone(user_tz)


def parse_ics(data: bytes, tz: str) -> list[Lesson]:
    user_tz = ZoneInfo(tz)
    seen: dict[tuple, Lesson] = {}
    for ev in _ics_events(data.decode("utf-8-sig", errors="replace")):
        if ev.get("STATUS", ([], ""))[1].strip().upper() == "CANCELLED" or "DTSTART" not in ev:
            continue
        start = _ics_datetime(*ev["DTSTART"], user_tz)
        if start is None:
            continue
        end = _ics_datetime(*ev["DTEND"], user_tz) if "DTEND" in ev else None
        if end is None:
            m = _DURATION.fullmatch(ev.get("DURATION", ([], ""))[1].strip())
            if m and any(m.groups()):
                d, h, mi, sec = (int(x or 0) for x in m.groups())
                end = start + timedelta(days=d, hours=h, minutes=mi, seconds=sec)
            else:
                end = start + timedelta(minutes=90)
        subject = _ics_unescape(ev.get("SUMMARY", ([], ""))[1]) or "Без названия"
        lesson = Lesson(
            weekday=start.weekday(),
            start=start.time().replace(second=0),
            end=end.time().replace(second=0),
            subject=subject[:100],
            room=_ics_unescape(ev.get("LOCATION", ([], ""))[1])[:100],
        )
        seen.setdefault((lesson.weekday, lesson.start, lesson.end, lesson.subject), lesson)
    if not seen:
        raise ScheduleParseError("в календаре не нашёл ни одного занятия со временем")
    return sorted(seen.values(), key=lambda l: (l.weekday, l.start))
